from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
SOURCE = BASE / "TOK_Uber_Eats_Level_Audit_2026-06-16.md"
OUT = BASE / "TOK_Uber_Eats_Level_Audit_2026-06-16.docx"
SCREENSHOTS = BASE / "screenshots-playwright"


INK = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(85, 85, 85)
ORANGE = RGBColor(255, 106, 26)
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_border_bottom(paragraph, color="DADCE0", size="8"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_masthead(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header_p = section.header.paragraphs[0]
    header_p.text = ""
    set_paragraph_spacing(header_p, 0, 0, 1)
    left = header_p.add_run("TOK Audit")
    set_run_font(left, size=9, color=MUTED, bold=True)
    header_p.add_run("  |  ")
    right = header_p.add_run("Niveau Uber Eats - 16 juin 2026")
    set_run_font(right, size=9, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer_p, 0, 0, 1)
    run = footer_p.add_run("Confidentiel - rapport d'audit technique et produit")
    set_run_font(run, size=8.5, color=MUTED)

    kicker = doc.add_paragraph()
    set_paragraph_spacing(kicker, 0, 2, 1.1)
    r = kicker.add_run("AUDIT TECHNIQUE, PRODUIT ET UX")
    set_run_font(r, size=10, color=ORANGE, bold=True)

    title = doc.add_paragraph()
    set_paragraph_spacing(title, 0, 4, 1.05)
    r = title.add_run("TOK vers niveau Uber Eats")
    set_run_font(r, size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, 0, 10, 1.2)
    r = subtitle.add_run(
        "Audit complet du code, des logiques backend, de Supabase, Stripe, GitHub, "
        "frontend et UX, avec roadmap de correction priorisee."
    )
    set_run_font(r, size=12, color=MUTED)

    meta_rows = [
        ("Repo", "mtnrconcept/cloud-rebuild"),
        ("Workspace", r"C:\Users\Pc\cloud-rebuild-recovered"),
        ("Branche", "main"),
        ("Date", "16 juin 2026"),
        ("Decision", "ne pas pousser production tant que les P0 ne sont pas corriges et verifies"),
    ]
    for label, value in meta_rows:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, 0, 2, 1.15)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, size=10.5, color=BLACK, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, size=10.5, color=BLACK)

    rule = doc.add_paragraph()
    add_border_bottom(rule, color="DADCE0", size="8")
    set_paragraph_spacing(rule, 2, 10, 1)


def inline_format(paragraph, text: str):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part[1:-1] if part.startswith("`") and part.endswith("`") else part)
        if part.startswith("`") and part.endswith("`"):
            set_run_font(run, name="Consolas", size=9.5, color=DARK_BLUE)
        else:
            set_run_font(run, size=11, color=BLACK)


def add_markdown_content(doc: Document, markdown: str):
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(line[3:].strip())
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(line[4:].strip())
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            p.add_run(line[5:].strip())
            continue
        if line == "---":
            p = doc.add_paragraph()
            add_border_bottom(p)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            inline_format(p, line[2:].strip())
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            inline_format(p, numbered.group(1).strip())
            continue
        p = doc.add_paragraph()
        set_paragraph_spacing(p, 0, 6, 1.25)
        inline_format(p, line)


def add_visual_appendix(doc: Document):
    assets = [
        ("Accueil desktop pendant l'intro", "home-1280x720.png"),
        ("Accueil desktop apres intro", "home-after-9s-1280x720.png"),
        ("Recherche desktop", "search-1280x720.png"),
        ("Accueil mobile apres intro", "home-mobile-after-9s-390x844.png"),
        ("Recherche mobile", "search-mobile-390x844.png"),
    ]
    doc.add_section(WD_SECTION.NEW_PAGE)
    h = doc.add_paragraph(style="Heading 1")
    h.add_run("Annexe visuelle")
    for caption, filename in assets:
        path = SCREENSHOTS / filename
        if not path.exists():
            continue
        p = doc.add_paragraph(style="Heading 2")
        p.add_run(caption)
        doc.add_picture(str(path), width=Inches(6.2))
        cap = doc.add_paragraph()
        set_paragraph_spacing(cap, 2, 10, 1.1)
        r = cap.add_run(filename)
        set_run_font(r, size=9, color=MUTED, italic=True)


def main():
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    add_masthead(doc)
    add_markdown_content(doc, markdown)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
