from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Pc\cloud-rebuild-recovered")
OUT_DIR = ROOT / "outputs" / "manual-20260607-tok-commercial-investor" / "documents" / "tok-sales-playbook"
FINAL_DOCX = OUT_DIR / "TOK_Playbook_Commercial_Restaurateurs.docx"
LOGO = ROOT / "public" / "logo.png"

TOK_ORANGE = "FF6A1A"
TOK_ORANGE_DARK = "D94F0F"
TOK_NAVY = "21314B"
TOK_INK = "0B1220"
TOK_MUTED = "64748B"
TOK_EMERALD = "16A874"
TOK_SKY = "22A3FF"
TOK_AMBER = "FFB02A"
TOK_WARM = "FFF2E8"
TOK_PAPER = "FFFDF9"
TOK_SOFT = "F4F6F9"
TOK_LINE = "E2E8F0"


def rgb(hex_color: str) -> RGBColor:
    value = hex_color.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_run_font(run, *, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color) if isinstance(color, str) else color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.strip("#"))


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_dxa = int(round(sum(widths_in) * 1440))
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    widths_dxa = [int(round(w * 1440)) for w in widths_in]
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_cell_text(cell, text, *, size=9.5, color=TOK_INK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    set_paragraph_spacing(p, after=0, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TOK_ORANGE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_body(doc: Document, text: str, *, bold=False, italic=False, color=TOK_INK, after=6, size=11):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after, line=1.25)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_kicker(doc: Document, text: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=2, line=1.0)
    run = p.add_run(text.upper())
    set_run_font(run, size=8.5, color=TOK_ORANGE, bold=True)
    return p


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_paragraph()
    style = doc.styles[f"Heading {level}"]
    p.style = style
    if p.runs:
        p.runs[0].text = ""
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, name="Georgia", size=16, color=TOK_NAVY, bold=True)
        set_paragraph_spacing(p, before=18, after=10, line=1.15)
    elif level == 2:
        set_run_font(run, size=13, color=TOK_NAVY, bold=True)
        set_paragraph_spacing(p, before=14, after=7, line=1.18)
    else:
        set_run_font(run, size=12, color=TOK_ORANGE_DARK, bold=True)
        set_paragraph_spacing(p, before=10, after=5, line=1.18)
    return p


def add_bullets(doc: Document, items, *, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        set_paragraph_spacing(p, after=4, line=1.25)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color=TOK_INK)


def add_numbered(doc: Document, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_spacing(p, after=4, line=1.25)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color=TOK_INK)


def add_callout(doc: Document, label: str, body: str, *, fill=TOK_WARM, accent=TOK_ORANGE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, start=220, end=220)
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=3, line=1.15)
    r = p.add_run(label.upper())
    set_run_font(r, size=8.5, color=accent, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.2)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=TOK_INK, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_table(doc: Document, headers, rows, widths, *, header_fill=TOK_NAVY, header_color="FFFFFF"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        shade_cell(hdr[idx], header_fill)
        set_cell_text(hdr[idx], header, size=8.8, color=header_color, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            if idx == 0:
                shade_cell(cells[idx], "FFF8F1")
                set_cell_text(cells[idx], str(value), size=9.3, color=TOK_NAVY, bold=True)
            else:
                set_cell_text(cells[idx], str(value), size=9.2, color=TOK_INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(TOK_INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for idx, size in [(1, 16), (2, 13), (3, 12)]:
        style = styles[f"Heading {idx}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(TOK_NAVY if idx < 3 else TOK_ORANGE_DARK)
        style.font.bold = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("TOK | Playbook commercial restaurateurs")
    set_run_font(r, size=8.5, color=TOK_MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Document interne - à adapter avec les chiffres terrain et les retours des commerciaux.")
    set_run_font(r, size=8.2, color=TOK_MUTED)


def cover(doc: Document):
    if LOGO.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_logo.add_run()
        run.add_picture(str(LOGO), width=Inches(1.4))
        set_paragraph_spacing(p_logo, before=4, after=8, line=1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=2, line=1.0)
    r = p.add_run("PLAYBOOK COMMERCIAL")
    set_run_font(r, name="Georgia", size=25, color=TOK_NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=12, line=1.1)
    r = p.add_run("Guide de vente restaurateurs - TOK / TheTok")
    set_run_font(r, size=13, color=TOK_ORANGE, bold=True)

    add_callout(
        doc,
        "Positionnement à tenir",
        "TOK se vend comme une plateforme restaurateur-first : plus de marge protégée, plus de relation directe, plus d'outils opérationnels et plus d'offres locales pour remplir les bons créneaux.",
    )

    table = add_table(
        doc,
        ["Signal", "Ce que le commercial doit retenir"],
        [
            ["Marché cible", "Restaurants suisses francophones, avec Genève comme zone d'attaque prioritaire."],
            ["Promesse B2B", "Remplir les tables, vendre les heures creuses, améliorer la visibilité et garder une lecture claire des coûts."],
            ["Preuves produit", "Dashboard restaurateur, packs de lancement, photos IA, offres anti-gaspi, ventes flash, campagnes, factures, plan de salle et support."],
            ["Modèle lisible", "5 CHF/table côté réservation, packs payants, campagnes sponsorisées, Tok One et modules IA selon activation."],
            ["Règle de vente", "Ne jamais promettre une automatisation non disponible en production ; vendre le plan d'activation et la trajectoire."],
        ],
        [1.55, 4.95],
        header_fill=TOK_ORANGE,
    )
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    add_body(doc, f"Version de travail - {date(2026, 6, 7).strftime('%d/%m/%Y')}", color=TOK_MUTED, size=9, after=0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def sales_strategy(doc: Document):
    add_heading(doc, "1. Thèse commerciale", 1)
    add_callout(
        doc,
        "La phrase d'ouverture",
        "Nous aidons les restaurants genevois à remplir les tables et créneaux creux sans dépendre d'une plateforme qui capte la marge et la relation client.",
    )
    add_body(
        doc,
        "Le commercial ne vend pas seulement une présence sur une app. Il vend un système : acquisition locale, conversion client, outils de pilotage, contenu premium et discipline financière.",
    )
    add_bullets(
        doc,
        [
            "Restaurateur indépendant : priorité au remplissage salle, à la marge et à la visibilité locale.",
            "Restaurant déjà présent sur des plateformes : priorité au coût, à la dépendance et à la relation client.",
            "Restaurant avec belle salle mais creux horaires : priorité aux réservations, offres ciblées, zéro attente et anti-gaspi.",
            "Restaurant avec mauvaise photo/menu : priorité au pack de lancement, menu digital, studio photo IA et campagnes.",
        ],
    )

    add_heading(doc, "2. Cibles prioritaires", 1)
    add_table(
        doc,
        ["Segment", "Douleur probable", "Angle de vente", "Premier produit à montrer"],
        [
            ["Bistrot / brasserie", "Creux en semaine, dépendance aux habitués", "Remplissage table + offres horaires", "Page B2B Genève + simulateur"],
            ["Cuisine rapide premium", "Visibilité, photos, panier moyen", "Photos IA + ventes flash + campagnes", "Dashboard photos et campagnes"],
            ["Restaurant gastronomique accessible", "Réservation, no-show, expérience", "Zéro attente + Table du Chef", "Parcours réservation prépayée"],
            ["Nouveau restaurant", "Lancement, carte, notoriété", "Pack Pro ou Premium", "Packs de lancement"],
            ["Restaurant déjà livré", "Commission et relation client captée", "Modèle plus lisible + CRM local", "Comparatif TOK vs plateformes classiques"],
        ],
        [1.35, 1.65, 1.85, 1.65],
    )

    add_heading(doc, "3. Qualification en 7 minutes", 1)
    add_numbered(
        doc,
        [
            "Quel service est le plus difficile à remplir aujourd'hui : midi, soir, semaine, dimanche ou fin de service ?",
            "Combien de tables couverts par mois voulez-vous ajouter sans casser l'équipe ?",
            "Quel est votre ticket moyen et votre marge ressentie sur les plateformes actuelles ?",
            "Avez-vous des photos de plats suffisamment bonnes pour convertir en ligne ?",
            "Quels canaux vous apportent vraiment des clients récurrents aujourd'hui ?",
            "Qui gère les réservations, les offres et les avis dans l'équipe ?",
            "Si TOK vous apporte X tables qualifiées ce mois-ci, quel serait le coût acceptable ?",
        ],
    )
    add_callout(
        doc,
        "Règle MEDDIC simplifiée",
        "Le deal est qualifié uniquement si le commercial identifie un besoin mesurable, un décideur, une contrainte opérationnelle et une prochaine étape datée.",
        fill="ECFDF5",
        accent=TOK_EMERALD,
    )
    doc.add_page_break()


def offer_architecture(doc: Document):
    add_heading(doc, "4. Architecture de l'offre", 1)
    add_body(doc, "L'offre se présente en escalier. Le but est d'éviter une vente trop large et de choisir un point d'entrée cohérent avec le besoin.")
    add_table(
        doc,
        ["Offre", "Prix / logique", "Quand la proposer", "Message court"],
        [
            ["Pack Découverte", "490 CHF - paiement unique", "Restaurant qui veut tester la plateforme proprement", "On met en place le socle et le menu digital."],
            ["Pack Essentiel", "990 CHF - paiement unique", "Restaurant qui manque de photos et présence digitale", "On rend le restaurant présentable et vendable."],
            ["Pack Pro", "1'990 CHF - populaire", "Restaurant prêt à lancer visibilité + campagnes", "On lance proprement : menu, photos, campagne, plan de salle."],
            ["Pack Premium", "3'490 CHF - VIP", "Restaurant à fort potentiel ou multi-sites", "On traite le lancement comme une vraie opération commerciale."],
            ["Réservations", "5 CHF/table confirmée", "Salle à remplir, créneaux à valoriser", "Coût lisible, lié à l'usage réel."],
            ["Campagnes", "CPM/CPC/CPA selon stratégie", "Besoin de trafic ou conversion", "On teste, mesure et réalloue."],
        ],
        [1.35, 1.35, 1.9, 1.9],
    )

    add_heading(doc, "5. Ce qu'il faut démo en priorité", 1)
    add_table(
        doc,
        ["Moment", "Écran / module", "Preuve à verbaliser"],
        [
            ["0-2 min", "Page restaurateurs Genève", "TOK a un discours B2B clair : marge, tables, photos et offres."],
            ["2-5 min", "Simulateur de marge", "Le restaurant comprend l'effet volume, coût TOK, marketing et marge."],
            ["5-8 min", "Packs de lancement", "L'accompagnement est tangible, pas une inscription vide."],
            ["8-12 min", "Dashboard restaurateur", "Le restaurateur garde le contrôle : menu, réservations, offres, factures, performances."],
            ["12-15 min", "Campagnes / Actualités", "TOK transforme le restaurant en média local activable."],
            ["15-18 min", "Plan de salle / service", "La promesse est opérationnelle, pas uniquement marketing."],
        ],
        [1.0, 1.75, 3.75],
        header_fill=TOK_NAVY,
    )

    add_callout(
        doc,
        "Ne pas sur-vendre",
        "Si le prospect demande une fonctionnalité encore fragile, répondre : 'Le module existe dans la trajectoire produit, mais on démarre sur ce qui est robuste et mesurable pour votre restaurant.'",
        fill="FFF7ED",
        accent=TOK_ORANGE_DARK,
    )
    doc.add_page_break()


def objection_handling(doc: Document):
    add_heading(doc, "6. Objections et réponses", 1)
    add_table(
        doc,
        ["Objection", "Réponse courte", "Question de relance"],
        [
            ["Je suis déjà sur Uber Eats / TheFork.", "Justement : TOK n'est pas seulement un canal, c'est un outil pour récupérer de la marge, des données et des créneaux.", "Aujourd'hui, quelle part de votre CA vient d'un canal que vous ne contrôlez pas ?"],
            ["Je n'ai pas le temps de gérer une plateforme de plus.", "Les packs existent pour absorber la mise en place et réduire l'effort initial.", "Qu'est-ce qui vous prend le plus de temps : menu, photos, réservations ou offres ?"],
            ["Je ne veux pas refaire des promos.", "TOK ne vend pas une remise permanente ; on active les bons créneaux et les bons stocks.", "Quels créneaux sont sous-remplis sans nuire à votre image ?"],
            ["Je veux voir des résultats avant de payer.", "On peut cadrer un plan d'activation avec jalons : menu publié, offres, premières tables, suivi hebdo.", "Quel résultat minimum rendrait l'essai sérieux pour vous ?"],
            ["Mes photos sont déjà bonnes.", "Très bien : on part alors sur réservation, campagnes, offres et pilotage plutôt que sur photo.", "Quel levier améliorerait le plus votre mois prochain ?"],
            ["Combien ça coûte vraiment ?", "Le modèle doit rester lisible : pack si accompagnement, 5 CHF/table côté réservation, campagnes mesurées.", "Préférez-vous maîtriser un coût fixe, un coût par table ou un budget campagne ?"],
        ],
        [1.75, 2.45, 2.3],
    )

    add_heading(doc, "7. Scripts prêts à utiliser", 1)
    add_callout(
        doc,
        "Appel à froid - 30 secondes",
        "Bonjour, je vous appelle de TOK. On aide les restaurants genevois à remplir leurs tables et leurs créneaux creux avec un modèle plus lisible que les plateformes classiques : réservation, offres locales, photos et pilotage. Je voulais voir si vous avez plutôt un sujet visibilité, réservations ou marge en ce moment.",
    )
    add_callout(
        doc,
        "Email après découverte",
        "Merci pour l'échange. J'ai retenu trois priorités : remplir [créneau], améliorer [levier] et garder un coût maîtrisé. Je vous propose une démo de 20 minutes centrée sur le simulateur de marge, les packs et le dashboard restaurateur, puis on décide ensemble du plan d'activation le plus simple.",
        fill="F8FAFC",
        accent=TOK_NAVY,
    )
    add_callout(
        doc,
        "Relance après silence",
        "Je vous relance car le sujet semblait clair : utiliser TOK pour convertir davantage de tables sans créer une mécanique de promos permanentes. Si ce n'est pas prioritaire ce mois-ci, je vous propose de reprendre quand vous préparez le prochain service fort ou la prochaine carte.",
        fill="ECFEFF",
        accent=TOK_SKY,
    )
    doc.add_page_break()


def pipeline(doc: Document):
    add_heading(doc, "8. Pipeline et discipline CRM", 1)
    add_table(
        doc,
        ["Statut", "Définition", "Sortie attendue"],
        [
            ["Nouveau", "Lead identifié mais non contacté.", "Appel ou email initial sous 24h."],
            ["Contacté", "Conversation engagée ou réponse reçue.", "Besoin principal noté + prochain créneau proposé."],
            ["Démo", "Démo planifiée ou réalisée.", "Décideur présent + offre recommandée."],
            ["Signé", "Accord commercial ou pack choisi.", "Paiement / onboarding / création restaurant."],
            ["Activé", "Menu publié, offre ou réservation active.", "Objectif 30/60/90 jours lancé."],
            ["Perdu", "Refus explicite ou absence de fit.", "Raison de perte et date de relance éventuelle."],
        ],
        [1.15, 3.05, 2.3],
        header_fill=TOK_ORANGE,
    )
    add_heading(doc, "9. Objectifs commerciaux", 1)
    add_bullets(
        doc,
        [
            "Objectif par commercial : 10 restaurants signés par mois comme base de pilotage interne.",
            "Commission commerciale de référence : 7 % dans la gouvernance locale, avec plafond à 12 % selon règles internes.",
            "Les commissions doivent suivre des jalons activables : pack payé, menu publié, premières tables, activité à 60/90 jours.",
            "Ne jamais pousser un pack trop élevé si le restaurant n'a pas l'équipe ou le volume pour l'activer.",
        ],
    )
    add_callout(
        doc,
        "Rituel hebdomadaire",
        "Chaque lundi : revoir nouveaux leads, démos planifiées, deals bloqués, raisons de perte, restaurants activés et premières métriques. Chaque vendredi : planifier les relances et les preuves à montrer la semaine suivante.",
        fill="F8FAFC",
        accent=TOK_NAVY,
    )

    add_heading(doc, "10. Score de priorité lead", 1)
    add_table(
        doc,
        ["Critère", "0 point", "1 point", "2 points"],
        [
            ["Besoin", "Vague", "Douleur claire", "Douleur chiffrée"],
            ["Décideur", "Inconnu", "Influenceur", "Propriétaire / gérant"],
            ["Urgence", "Aucune", "Sous 3 mois", "Sous 30 jours"],
            ["Fit TOK", "Hors cible", "Un module utile", "Packs + réservations + offres"],
            ["Capacité d'activation", "Équipe bloquée", "Possible mais lent", "Prêt à publier / tester"],
        ],
        [1.3, 1.7, 1.7, 1.8],
    )
    add_body(doc, "Règle : 7 points ou plus = lead prioritaire ; 5-6 = nurturing ; moins de 5 = relance légère ou no-fit.", bold=True, color=TOK_NAVY)
    doc.add_page_break()


def field_enablement(doc: Document):
    add_heading(doc, "11. Checklists terrain", 1)
    add_heading(doc, "Avant la démo", 2)
    add_bullets(
        doc,
        [
            "Identifier le type de restaurant, son quartier, ses créneaux faibles et son ticket moyen approximatif.",
            "Préparer une hypothèse de pack : Découverte, Essentiel, Pro ou Premium.",
            "Préparer une démo courte : simulateur de marge, packs, dashboard et un levier métier seulement.",
            "Noter les objections probables : commission, temps, visibilité, concurrence, coût initial.",
        ],
    )
    add_heading(doc, "Pendant la démo", 2)
    add_bullets(
        doc,
        [
            "Commencer par la douleur du restaurateur, pas par les fonctionnalités.",
            "Faire manipuler les chiffres : tables/mois, ticket moyen, coût marketing.",
            "Choisir un plan d'activation en une phrase : 'On commence par...'.",
            "Terminer avec une prochaine étape datée : pack, compte, rendez-vous, envoi des photos ou carte.",
        ],
    )
    add_heading(doc, "Après la signature", 2)
    add_bullets(
        doc,
        [
            "Confirmer le pack, le responsable, les assets nécessaires et le calendrier.",
            "Suivre menu publié, photos prêtes, première offre, premières réservations et premiers retours.",
            "Ne pas laisser le restaurant signé sans activation : la valeur doit apparaître dans les 14 premiers jours.",
        ],
    )

    add_heading(doc, "12. Messages à ne jamais envoyer", 1)
    add_table(
        doc,
        ["À éviter", "Pourquoi", "Alternative"],
        [
            ["On va vous ramener beaucoup de clients.", "Promesse vague et risquée.", "On cadre un objectif mesurable de tables, offres ou visibilité."],
            ["C'est comme Uber Eats mais moins cher.", "Positionnement défensif.", "TOK combine réservation, offres, relation client et outils restaurateur."],
            ["Tout est automatisé.", "Peut créer une attente fausse.", "Les modules clés sont outillés, avec accompagnement selon pack."],
            ["Aucun effort de votre côté.", "Faux pour menu, photos, validation.", "On réduit fortement l'effort de mise en place."],
        ],
        [1.6, 2.2, 2.7],
        header_fill=TOK_NAVY,
    )
    doc.add_page_break()


def appendix(doc: Document):
    add_heading(doc, "13. Sources internes et éléments à personnaliser", 1)
    add_body(doc, "Ce playbook est basé sur l'état local du repo TOK au 7 juin 2026. Avant diffusion externe, compléter avec les métriques réelles terrain : restaurants signés, taux de démo, activation, tables générées, rétention et revenus.")
    add_table(
        doc,
        ["Source", "Usage dans le playbook"],
        [
            ["src/pages/RestaurateursGeneve.tsx", "Message B2B Genève, simulateur de marge, 5 CHF/table, packs mensuels d'exemple."],
            ["supabase/migrations/20260404120000_launch_offer_packs.sql", "Prix et contenu des packs Découverte, Essentiel, Pro, Premium."],
            ["supabase/migrations/20260607033000_platform_finance_sales_governance.sql", "Pipeline leads/deals, objectif 10 restaurants/mois, commission commerciale 7 %."],
            ["docs/skills/TOK_APPLICATION_SKILL.md", "Positionnement restaurateur-first, modules client/dashboard/courier/admin."],
            ["docs/audits/2026-05-31-global-site-gap-audit.md", "Maturité plateforme et priorités opérationnelles."],
        ],
        [2.35, 4.15],
        header_fill=TOK_ORANGE,
    )
    add_heading(doc, "Liens utiles pour contexte marché", 2)
    p = add_body(doc, "Population suisse : ", after=2)
    add_hyperlink(p, "communiqué FSO / geo.admin.ch", "https://www.geo.admin.ch/en/nsb?id=104703")
    p = add_body(doc, "Population du canton de Genève : ", after=2)
    add_hyperlink(p, "OCSTAT Genève", "https://statistique.ge.ch/domaines/apercu.asp?dom=01_01")
    p = add_body(doc, "Définition statistique de la restauration : ", after=2)
    add_hyperlink(p, "NOGA 2025 - activités de service de restauration", "https://www.kubb-tool.bfs.admin.ch/fr/noga/2025/56")

    add_callout(
        doc,
        "À compléter par l'équipe",
        "Ajouter la grille tarifaire validée, les preuves clients réelles, le CRM cible, les règles de commission définitives et la politique de remise autorisée.",
        fill="FFF7ED",
        accent=TOK_ORANGE_DARK,
    )


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    cover(doc)
    sales_strategy(doc)
    offer_architecture(doc)
    objection_handling(doc)
    pipeline(doc)
    field_enablement(doc)
    appendix(doc)
    doc.save(FINAL_DOCX)
    print(FINAL_DOCX)


if __name__ == "__main__":
    build()
