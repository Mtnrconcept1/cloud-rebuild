from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Pc\cloud-rebuild-recovered")
OUT_DIR = ROOT / "outputs" / "manual-20260607-tok-commercial-investor" / "documents" / "tok-sales-playbook-v2"
FINAL_DOCX = OUT_DIR / "TOK_Playbook_Commercial_Restaurateurs_V2_Digest.docx"

LOGO = ROOT / "public" / "logo.png"
FUN = ROOT / "public" / "fond fun.png"
MOCKUP = ROOT / "public" / "mockup.png"

ORANGE = "FF6A1A"
ORANGE_DARK = "D94F0F"
NAVY = "0B1220"
BLUE = "21314B"
PAPER = "FFF8F1"
WARM = "FFF2E8"
CREAM = "FFFDF8"
MINT = "ECFDF5"
SKY = "EAF6FF"
YELLOW = "FFF7D6"
PINK = "FFF1F2"
LINE = "F0D6C2"
MUTED = "64748B"
EMERALD = "16A874"
CYAN = "22A3FF"
AMBER = "FFB02A"


def rgb(value: str) -> RGBColor:
    value = value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def font(run, *, name="Aptos", size=10.5, color=NAVY, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def spacing(paragraph, *, before=0, after=5, line=1.15):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=120, start=160, bottom=120, end=160):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    total = int(round(sum(widths) * 1440))
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    widths_dxa = [int(round(width * 1440)) for width in widths]
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def cell_text(cell, lines, *, size=10, color=NAVY, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    if isinstance(lines, str):
        lines = [lines]
    for idx, text in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = align
        spacing(p, after=2, line=1.1)
        run = p.add_run(text)
        font(run, size=size, color=color, bold=bold if idx == 0 else False)


def para(doc, text="", *, size=10.5, color=NAVY, bold=False, italic=False, after=6, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    spacing(p, after=after, line=1.18)
    r = p.add_run(text)
    font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def heading(doc, text, *, level=1, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        spacing(p, before=8, after=2, line=1)
        r = p.add_run(kicker.upper())
        font(r, size=8.5, color=ORANGE, bold=True)
    p = doc.add_paragraph()
    spacing(p, before=8 if level == 1 else 4, after=6, line=1.05)
    r = p.add_run(text)
    font(r, name="Georgia", size=18 if level == 1 else 13.5, color=BLUE, bold=True)
    return p


def card_row(doc, cards, fills=None):
    fills = fills or [WARM, MINT, SKY]
    table = doc.add_table(rows=1, cols=len(cards))
    table.style = "Table Grid"
    table_geometry(table, [6.5 / len(cards)] * len(cards))
    for idx, (title, body) in enumerate(cards):
        cell = table.cell(0, idx)
        shade(cell, fills[idx % len(fills)])
        cell_text(cell, [title, body], size=9.5 if len(cards) > 3 else 10.2, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(7)


def ticket(doc, label, title, body, *, fill=WARM, accent=ORANGE):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table_geometry(table, [1.05, 5.45])
    shade(table.cell(0, 0), accent)
    shade(table.cell(0, 1), fill)
    cell_text(table.cell(0, 0), label.upper(), size=10, color="FFFFFF", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(table.cell(0, 1), [title, body], size=10.4, color=NAVY, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(5)


def quick_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table_geometry(table, widths)
    for idx, header in enumerate(headers):
        shade(table.cell(0, idx), BLUE)
        cell_text(table.cell(0, idx), header, size=8.8, color="FFFFFF", bold=True)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            shade(cells[idx], CREAM if row_idx % 2 == 0 else "FFFFFF")
            cell_text(cells[idx], str(value), size=9.0, color=NAVY, bold=idx == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    for style_name in ["Normal", "Body Text"]:
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = rgb(NAVY)
    header = section.header.paragraphs[0]
    header.text = "TOK - Playbook commercial digest"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if header.runs:
        font(header.runs[0], size=8.5, color=MUTED, bold=True)


def cover(doc):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.0))
    para(doc, "PLAYBOOK COMMERCIAL", size=9, color=ORANGE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, after=8, line=1.02)
    r = p.add_run("La bonne table,\npas la blabla.")
    font(r, name="Georgia", size=28, color=BLUE, bold=True)
    para(
        doc,
        "Guide terrain pour vendre TOK aux restaurateurs : plus clair, plus vivant, plus facile à sortir en rendez-vous.",
        size=12,
        color=NAVY,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    if FUN.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FUN), width=Inches(4.9))
    card_row(
        doc,
        [
            ("Objectif", "Remplir les bons créneaux."),
            ("Promesse", "Moins de commission floue, plus de contrôle."),
            ("Ton", "Direct, premium, souriant."),
        ],
        [WARM, MINT, SKY],
    )
    para(doc, "Version V2 digest - 7 juin 2026", size=8.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def core(doc):
    heading(doc, "1. Le pitch qui tient entre deux services", kicker="À dire simplement")
    ticket(
        doc,
        "15 sec",
        "TOK aide les restaurants genevois à remplir leurs tables et leurs créneaux utiles.",
        "On combine réservation, offres locales, photos, dashboard et pilotage. Pas une vitrine de plus : un outil pour vendre mieux.",
        fill=WARM,
    )
    card_row(
        doc,
        [
            ("Ce qu'on ne vend pas", "Une app magique ou une promo permanente."),
            ("Ce qu'on vend", "Un plan d'activation restaurant par restaurant."),
            ("Le bon mot", "Marge, tables, contrôle, simplicité."),
        ],
        [PINK, MINT, YELLOW],
    )

    heading(doc, "2. Les prospects croustillants", kicker="Priorité terrain")
    quick_table(
        doc,
        ["Cible", "Douleur visible", "Angle TOK"],
        [
            ("Salle irrégulière", "Mardi midi vide, dimanche mou, fin de service molle.", "Réservation + offres ciblées."),
            ("Belle carte, peu de digital", "Plats forts mais photos / page peu vendeuses.", "Pack Essentiel ou Pro."),
            ("Déjà livré ailleurs", "Marge compressée, relation client captée.", "Modèle plus lisible + dashboard."),
            ("Nouveau lieu", "Besoin de notoriété et d'un lancement propre.", "Pack Pro / Premium."),
        ],
        [1.55, 2.55, 2.4],
    )
    ticket(
        doc,
        "Score",
        "7/10 ou plus = on fonce.",
        "Besoin clair, décideur identifié, urgence sous 30 jours, capacité à activer vite, fit TOK visible.",
        fill=MINT,
        accent=EMERALD,
    )

    heading(doc, "3. La démo en mode menu dégustation", kicker="Rythme recommandé")
    quick_table(
        doc,
        ["Minute", "Plat à servir", "Phrase courte"],
        [
            ("0-2", "Douleur du resto", "On part de votre service à remplir."),
            ("2-5", "Simulateur", "Voilà ce qu'une table change vraiment."),
            ("5-9", "Packs", "On choisit l'effort d'activation juste."),
            ("9-14", "Dashboard", "Vous gardez la main sur menu, offres, avis, factures."),
            ("14-18", "Next step", "On bloque le prochain service à améliorer."),
        ],
        [0.8, 2.1, 3.6],
    )
    doc.add_page_break()


def offer(doc):
    heading(doc, "4. L'offre sans prise de tête", kicker="Prix lisibles")
    card_row(
        doc,
        [
            ("Découverte\n490 CHF", "Socle, compte, menu, démarrage propre."),
            ("Essentiel\n990 CHF", "Photos, présence digitale, premières offres."),
            ("Pro\n1'990 CHF", "Lancement complet, campagne, plan de salle."),
            ("Premium\n3'490 CHF", "VIP, multi-leviers, accompagnement renforcé."),
        ],
        [WARM, SKY, YELLOW, MINT],
    )
    card_row(
        doc,
        [
            ("5 CHF/table", "Réservation confirmée, coût lisible."),
            ("Campagnes", "CPM/CPC/CPA selon objectif."),
            ("Tok One", "Récurrence client et avantages."),
        ],
        [MINT, SKY, YELLOW],
    )
    ticket(
        doc,
        "Règle",
        "On ne vend pas tout le buffet au premier rendez-vous.",
        "On choisit le levier qui règle le vrai problème du restaurateur : table, visibilité, marge, contenu ou pilotage.",
        fill=PAPER,
    )

    heading(doc, "5. Objections qui passent crème", kicker="Réponses courtes")
    quick_table(
        doc,
        ["Objection", "Réponse TOK", "Relance"],
        [
            ("Je suis déjà sur d'autres plateformes.", "Justement : TOK remet la marge et la relation client dans vos mains.", "Quel canal vous coûte le plus cher aujourd'hui ?"),
            ("Je n'ai pas le temps.", "Les packs absorbent la mise en place. Vous validez, on structure.", "Qu'est-ce qui vous prend le plus de temps ?"),
            ("Je ne veux pas faire des promos.", "On active les bons créneaux, pas une braderie permanente.", "Quel service mérite un coup de pouce ?"),
            ("Je veux voir avant de payer.", "On peut cadrer un plan clair : menu publié, offre lancée, premières tables suivies.", "Quel résultat rendrait l'essai sérieux ?"),
            ("Combien ça coûte vraiment ?", "Pack si accompagnement, 5 CHF/table côté réservation, campagnes mesurées.", "Vous préférez coût fixe, par table ou budget campagne ?"),
        ],
        [1.55, 2.55, 2.4],
    )
    doc.add_page_break()


def scripts(doc):
    heading(doc, "6. Scripts prêts à sortir", kicker="Copier, adapter, envoyer")
    ticket(
        doc,
        "Appel",
        "Bonjour, je vous appelle de TOK.",
        "On aide les restaurants genevois à remplir leurs tables et leurs créneaux creux avec un modèle plus lisible : réservation, offres locales, photos et pilotage. Aujourd'hui, votre sujet principal, c'est plutôt visibilité, marge ou réservations ?",
        fill=WARM,
    )
    ticket(
        doc,
        "Email",
        "Suite à notre échange, j'ai retenu trois priorités.",
        "Remplir [créneau], améliorer [levier] et garder un coût maîtrisé. Je vous propose une démo de 20 minutes : simulateur, packs, dashboard, puis décision sur le plan le plus simple.",
        fill=SKY,
        accent=CYAN,
    )
    ticket(
        doc,
        "Relance",
        "Je vous relance sans vous rajouter un service de plus.",
        "Si le sujet n'est pas prioritaire ce mois-ci, on reprend au moment de votre prochain temps fort : nouvelle carte, terrasse, service du dimanche ou période creuse.",
        fill=MINT,
        accent=EMERALD,
    )

    heading(doc, "7. Discipline CRM : simple mais carrée", kicker="Pas de sauce floue")
    quick_table(
        doc,
        ["Statut", "Définition", "Sortie attendue"],
        [
            ("Nouveau", "Lead identifié.", "Contact sous 24h."),
            ("Contacté", "Réponse ou conversation.", "Douleur + prochaine date."),
            ("Démo", "Créneau fixé ou démo faite.", "Décideur + offre recommandée."),
            ("Signé", "Pack ou accord choisi.", "Paiement / onboarding."),
            ("Activé", "Menu, offre ou réservation live.", "Objectif 30/60/90 jours."),
            ("Perdu", "No-fit ou refus.", "Raison claire + relance éventuelle."),
        ],
        [1.2, 2.55, 2.75],
    )
    ticket(
        doc,
        "Rituel",
        "Lundi on choisit les tables à gagner. Vendredi on vérifie ce qui a bougé.",
        "Nouveaux leads, démos, deals bloqués, raisons de perte, restaurants activés et preuves terrain.",
        fill=YELLOW,
        accent=AMBER,
    )
    doc.add_page_break()


def close(doc):
    heading(doc, "8. La checklist avant rendez-vous", kicker="À garder sous la main")
    card_row(
        doc,
        [
            ("Avant", "Quartier, type de resto, créneau faible, ticket moyen, pack probable."),
            ("Pendant", "Partir de la douleur. Faire parler les chiffres. Choisir un seul levier."),
            ("Après", "Prochaine date, assets nécessaires, pack, responsable, activation J+14."),
        ],
        [WARM, SKY, MINT],
    )
    heading(doc, "9. Phrases interdites", kicker="À éviter")
    quick_table(
        doc,
        ["Ne pas dire", "Dire plutôt"],
        [
            ("On va vous ramener beaucoup de clients.", "On cadre un objectif mesurable de tables, offres ou visibilité."),
            ("C'est comme les plateformes, mais moins cher.", "TOK combine réservation, offres, relation client et outils restaurateur."),
            ("Tout est automatisé.", "On outille les modules clés et on accompagne selon le pack."),
            ("Aucun effort de votre côté.", "On réduit l'effort, puis on active vite avec votre validation."),
        ],
        [2.75, 3.75],
    )
    heading(doc, "10. Sources et données à compléter", kicker="Avant diffusion externe")
    para(doc, "Base locale utilisée : modules restaurateurs Genève, packs de lancement, gouvernance commerciale, audit produit et contexte TOK. Compléter avant envoi externe avec les preuves réelles : restaurants signés, taux de démo, activation J+14, tables générées, CAC, churn, revenus et marge.", size=9.7, color=NAVY)
    card_row(
        doc,
        [
            ("À garder", "Ton direct, premium, souriant."),
            ("À ajouter", "Cas clients, screenshots validés, chiffres terrain."),
            ("À mesurer", "Démo -> signé -> activé -> tables."),
        ],
        [MINT, YELLOW, SKY],
    )


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    cover(doc)
    core(doc)
    offer(doc)
    scripts(doc)
    close(doc)
    doc.save(FINAL_DOCX)
    print(FINAL_DOCX)


if __name__ == "__main__":
    build()
